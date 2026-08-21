"""Investor spreadsheet import. Invalid rows are reported, never silently dropped."""

import csv
import re
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.investor import Investor
from app.models.investor_criteria import DEFAULT_STRICT_FIELDS, InvestorCriteria
from app.models.realtor import Realtor
from app.services.audit import AuditService
from app.services.matching.screening import normalize_property_type
from app.utilities.ids import next_public_id
from app.utilities.money import as_decimal
from app.utilities.parsing import as_bool, as_float, as_int, first_present, split_list

COLUMN_ALIASES = {
    "name": ["investor name", "name", "company", "entity"],
    "contact_name": ["contact name", "contact", "primary contact"],
    "phone": ["phone", "mobile", "cell"],
    "email": ["email", "e-mail"],
    "channel": ["preferred channel", "channel", "preferred notification method"],
    "may_text": ["may we text?", "may we text", "text ok", "sms ok"],
    "may_email": ["may we email?", "may we email", "email ok"],
    "min_price": ["minimum price", "min price", "min_price"],
    "max_price": ["maximum price", "max price", "max_price"],
    "cities": ["buying areas", "cities", "city", "coverage area"],
    "zip_codes": ["zip codes", "zips", "zip", "zip code"],
    "property_types": ["property types", "property type", "types"],
    "bedrooms": ["bedrooms", "beds", "min bedrooms"],
    "bathrooms": ["bathrooms", "baths", "min bathrooms"],
    "min_sqft": ["minimum sq ft", "min sq ft", "min_sqft", "sqft"],
    "max_hoa": ["max hoa", "max_hoa", "max hoa $"],
    "hoa": ["hoa allowed?", "hoa required", "hoa"],
    "financing": ["cash / finance", "financing", "cash/finance", "funding"],
    "notes": ["assigned notes", "notes", "comments"],
    "profile_name": ["profile", "criteria profile", "profile name"],
    "year_built_min": ["year built min", "min year built", "year_built_min"],
    "max_price_pct_of_arv": [
        "maximum purchase price (% arv)",
        "max purchase price (% arv)",
        "max price % arv",
        "% arv",
        "max_price_pct_of_arv",
    ],
    "min_desired_discount": ["required arv discount", "desired discount", "min desired discount"],
}

CITY_ALIASES = {
    "lv": "Las Vegas",
    "las vegas": "Las Vegas",
    "nlv": "North Las Vegas",
    "n las vegas": "North Las Vegas",
    "n. las vegas": "North Las Vegas",
    "north las vegas": "North Las Vegas",
    "henderson": "Henderson",
}


class InvestorImportService:
    def __init__(self, db: Session, audit: AuditService | None = None) -> None:
        self.db = db
        self.audit = audit or AuditService(db)

    def import_path(self, realtor: Realtor, path: Path, source: str | None = None) -> dict:
        rows = _read_tabular(path)
        return self.import_rows(realtor, rows, source=source or path.name)

    def import_bytes(self, realtor: Realtor, data: bytes, filename: str) -> dict:
        suffix = Path(filename).suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            rows = _read_xlsx(data)
        else:
            text = data.decode("utf-8-sig")
            rows = list(csv.DictReader(StringIO(text)))
        return self.import_rows(realtor, rows, source=filename)

    def import_rows(self, realtor: Realtor, rows: list[dict], source: str) -> dict:
        created_investors = 0
        created_profiles = 0
        errors: list[dict] = []
        for index, raw in enumerate(rows, start=2):
            if not any(str(v).strip() for v in raw.values() if v is not None):
                continue
            try:
                payload = _normalize_row(raw)
                investor = (
                    self.db.query(Investor)
                    .filter(Investor.realtor_id == realtor.id, Investor.name == payload["name"])
                    .one_or_none()
                )
                if investor is None:
                    investor = Investor(
                        public_id=next_public_id(self.db, "INV"),
                        realtor_id=realtor.id,
                        name=payload["name"],
                        contact_name=payload.get("contact_name"),
                        phone=payload.get("phone"),
                        email=payload.get("email"),
                        preferred_channel=payload.get("channel") or "none",
                        notes=payload.get("notes"),
                        communication_permissions=payload["permissions"],
                        import_source=source,
                    )
                    self.db.add(investor)
                    self.db.flush()
                    created_investors += 1
                else:
                    if payload.get("contact_name"):
                        investor.contact_name = payload["contact_name"]
                    if payload.get("phone"):
                        investor.phone = payload["phone"]
                    if payload.get("email"):
                        investor.email = payload["email"]
                    if payload.get("channel"):
                        investor.preferred_channel = payload["channel"]
                    investor.communication_permissions = payload["permissions"]
                    if payload.get("notes"):
                        investor.notes = payload["notes"]
                    investor.import_source = source
                    self.db.flush()
                profile_name = payload.get("profile_name") or "Imported profile"
                criteria = (
                    self.db.query(InvestorCriteria)
                    .filter(
                        InvestorCriteria.investor_id == investor.id,
                        InvestorCriteria.name == profile_name,
                    )
                    .one_or_none()
                )
                if criteria is None:
                    criteria = InvestorCriteria(
                        public_id=next_public_id(self.db, "CRT"),
                        realtor_id=realtor.id,
                        investor_id=investor.id,
                        name=profile_name,
                        strict_fields=list(DEFAULT_STRICT_FIELDS),
                    )
                    self.db.add(criteria)
                    self.db.flush()
                    created_profiles += 1
                criteria.min_price = payload.get("min_price")
                criteria.max_price = payload.get("max_price")
                criteria.cities = payload.get("cities") or []
                criteria.zip_codes = payload.get("zip_codes") or []
                criteria.property_types = payload.get("property_types") or []
                criteria.min_bedrooms = payload.get("bedrooms")
                criteria.min_bathrooms = payload.get("bathrooms")
                criteria.min_sqft = payload.get("min_sqft")
                criteria.max_hoa = payload.get("max_hoa")
                criteria.hoa_required = payload.get("hoa_required")
                criteria.year_built_min = payload.get("year_built_min")
                criteria.max_price_pct_of_arv = payload.get("max_price_pct_of_arv")
                criteria.min_desired_discount = payload.get("min_desired_discount")
                criteria.preferred_financing = payload.get("financing")
                criteria.notes = payload.get("notes")
                self.db.flush()
            except Exception as exc:  # noqa: BLE001
                errors.append({"row": index, "error": str(exc), "data": {k: str(v) for k, v in raw.items()}})
        self.audit.record(
            event="INVESTORS_IMPORTED",
            object_type="investor_import",
            object_id=source,
            actor="import_service",
            realtor_id=str(realtor.id),
            after_state={
                "created_investors": created_investors,
                "created_profiles": created_profiles,
                "errors": len(errors),
            },
        )
        return {
            "created_investors": created_investors,
            "created_profiles": created_profiles,
            "errors": errors,
            "rows_processed": len(rows),
        }


def _normalize_row(raw: dict) -> dict:
    name = first_present(raw, COLUMN_ALIASES["name"])
    if not name:
        raise ValueError("Investor Name is required")
    channel = _normalize_channel(first_present(raw, COLUMN_ALIASES["channel"]))
    types = [
        normalize_property_type(item)
        for item in _split_areas(first_present(raw, COLUMN_ALIASES["property_types"]))
        if item
    ]
    hoa_required, max_hoa = _parse_hoa(
        first_present(raw, COLUMN_ALIASES["hoa"]),
        first_present(raw, COLUMN_ALIASES["max_hoa"]),
    )
    financing = _normalize_financing(first_present(raw, COLUMN_ALIASES["financing"]))
    may_text = as_bool(first_present(raw, COLUMN_ALIASES["may_text"]))
    may_email = as_bool(first_present(raw, COLUMN_ALIASES["may_email"]))
    return {
        "name": str(name).strip(),
        "contact_name": _str(first_present(raw, COLUMN_ALIASES["contact_name"])),
        "phone": _str(first_present(raw, COLUMN_ALIASES["phone"])),
        "email": _str(first_present(raw, COLUMN_ALIASES["email"])),
        "channel": channel,
        "permissions": {"sms": may_text is True, "email": may_email is True},
        "min_price": as_decimal(first_present(raw, COLUMN_ALIASES["min_price"])),
        "max_price": as_decimal(first_present(raw, COLUMN_ALIASES["max_price"])),
        "cities": _normalize_cities(first_present(raw, COLUMN_ALIASES["cities"])),
        "zip_codes": split_list(first_present(raw, COLUMN_ALIASES["zip_codes"])),
        "property_types": types,
        "bedrooms": as_int(first_present(raw, COLUMN_ALIASES["bedrooms"])),
        "bathrooms": as_float(first_present(raw, COLUMN_ALIASES["bathrooms"])),
        "min_sqft": as_int(first_present(raw, COLUMN_ALIASES["min_sqft"])),
        "max_hoa": max_hoa,
        "hoa_required": hoa_required,
        "notes": _str(first_present(raw, COLUMN_ALIASES["notes"])),
        "profile_name": _str(first_present(raw, COLUMN_ALIASES["profile_name"])),
        "financing": financing,
        "year_built_min": as_int(first_present(raw, COLUMN_ALIASES["year_built_min"])),
        "max_price_pct_of_arv": _as_percent(first_present(raw, COLUMN_ALIASES["max_price_pct_of_arv"])),
        "min_desired_discount": _as_percent(first_present(raw, COLUMN_ALIASES["min_desired_discount"])),
    }


def _str(value: object | None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _read_tabular(path: Path) -> list[dict]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _read_xlsx(path.read_bytes())
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _read_xlsx(data: bytes) -> list[dict]:
    workbook = load_workbook(filename=BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header_index = _header_row_index(rows)
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[header_index]]
    result = []
    for row in rows[header_index + 1 :]:
        if not any(cell not in (None, "") for cell in row):
            continue
        result.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return result


def _header_row_index(rows: list[tuple]) -> int:
    for index, row in enumerate(rows[:20]):
        cells = [str(cell).strip().lower() if cell is not None else "" for cell in row]
        blob = " ".join(cells)
        if "investor name" in blob:
            return index
        if "name" in cells and ("phone" in cells or "email" in cells):
            return index
    return 0


def _normalize_channel(value: object | None) -> str | None:
    text = _str(value)
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"text", "mobile", "sms"}:
        return "sms"
    if lowered in {"phone call", "call", "phone", "voice"}:
        return "phone"
    if lowered in {"email", "e-mail"}:
        return "email"
    return lowered


def _normalize_financing(value: object | None) -> str | None:
    text = _str(value)
    if not text:
        return None
    lowered = text.lower()
    if "cash" in lowered and "financ" not in lowered:
        return "cash"
    return text


def _normalize_cities(value: object | None) -> list[str]:
    found: list[str] = []
    for item in _split_areas(value):
        key = re.sub(r"\s+", " ", item.lower()).strip()
        found.append(CITY_ALIASES.get(key, item))
    return found


def _split_areas(value: object | None) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        parts = value
    else:
        parts = re.split(r"[,;/|]+", str(value))
    return [str(part).strip() for part in parts if str(part).strip()]


def _parse_hoa(hoa_value: object | None, max_hoa_value: object | None) -> tuple[bool | None, Decimal | None]:
    max_hoa = as_decimal(max_hoa_value)
    raw = _str(hoa_value)
    if not raw:
        return None, max_hoa
    lowered = raw.lower()
    if "no hoa" in lowered or lowered in {"no", "false", "none"}:
        return False, max_hoa
    parsed = as_bool(raw)
    if parsed is not None:
        return parsed, max_hoa
    numeric = as_decimal(raw)
    if numeric is not None:
        return None, numeric
    return None, max_hoa


def _as_percent(value: object | None) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        number = value
    else:
        cleaned = str(value).replace("%", "").replace(",", "").strip()
        parsed = as_float(cleaned)
        if parsed is None:
            return None
        number = Decimal(str(parsed))
    if number > 1:
        number = (number / Decimal("100")).quantize(Decimal("0.0001"))
    return number
