from wellness_agent.catalog import UnknownProductError, find_product, format_catalog, pdf_path, products
from wellness_agent.cli import main
from wellness_agent.telegram_inbound import handle_telegram_update


class _FakeTelegram:
    def __init__(self) -> None:
        self.sent = []

    def send_message(self, chat_id, text, reply_markup=None):
        self.sent.append({"chat_id": chat_id, "text": text, "reply_markup": reply_markup})
        return {"ok": True}

    def send_document(self, chat_id, path, caption="", reply_markup=None, filename=None):
        self.sent.append(
            {
                "chat_id": chat_id,
                "document": str(path),
                "caption": caption,
                "reply_markup": reply_markup,
                "filename": filename,
            }
        )
        return {"ok": True}

    def send_photo(self, chat_id, path, caption="", reply_markup=None):
        self.sent.append({"chat_id": chat_id, "photo": str(path), "caption": caption, "reply_markup": reply_markup})
        return {"ok": True}

    def answer_callback_query(self, callback_query_id, text=None):
        return {"ok": True}


def test_live_shop_inventory_has_eight_skus_with_pdfs():
    items = products()
    ids = [item["id"] for item in items]
    assert ids == [
        "tirzepatide",
        "retatrutide",
        "semax",
        "nad",
        "klow",
        "mots-c",
        "ss-31",
        "ghk-cu",
    ]
    for item in items:
        path = pdf_path(item)
        assert path.is_file()
        assert path.suffix == ".pdf"
        assert path.stat().st_size > 1000


def test_find_product_aliases():
    assert find_product("tirz")["id"] == "tirzepatide"
    assert find_product("NAD+")["id"] == "nad"
    assert find_product("reta")["id"] == "retatrutide"
    assert find_product("GHK-Cu")["id"] == "ghk-cu"
    assert find_product("ss31")["id"] == "ss-31"


def test_unknown_product_is_rejected():
    try:
        find_product("hormuz-catalyst")
        raise AssertionError("expected UnknownProductError")
    except UnknownProductError as exc:
        assert "catalog" in str(exc)


def test_catalog_cli(capsys):
    assert main(["catalog"]) == 0
    out = capsys.readouterr().out
    assert "KLOW" in out
    assert "MOTS-c" in out
    assert "SS-31" in out
    assert "GHK-Cu" in out
    assert "Tirzepatide" in out


def test_product_cli(capsys):
    assert main(["product", "klow"]) == 0
    assert "klow.pdf" in capsys.readouterr().out


def test_telegram_catalog_and_product_sheet():
    tg = _FakeTelegram()
    catalog = handle_telegram_update(
        {"message": {"text": "/catalog", "chat": {"id": 7}, "from": {"id": 7}}},
        tg,
    )
    assert catalog["action"] == "menu"
    product = handle_telegram_update(
        {"message": {"text": "/product klow", "chat": {"id": 7}, "from": {"id": 7}}},
        tg,
    )
    assert product == {"ok": True, "action": "product", "chat_id": "7", "product": "klow"}
    docs = [item for item in tg.sent if "document" in item]
    assert docs
    assert docs[0]["document"].endswith("klow.pdf")
    assert "Educational only" in docs[0]["caption"]
    assert "Tap it to view" in docs[0]["caption"]
    assert "trueholdwellness.com/ols/" not in docs[0]["caption"]
    assert "Zelle" in docs[0]["caption"]


def test_all_sheets_are_vial_specific_and_segmented():
    from pypdf import PdfReader

    from wellness_agent.inventory.build_pdfs import build_missing
    from wellness_agent.inventory.protocol import PROTOCOLS

    build_missing()
    required = (
        "what it is",
        "how it works in the body",
        "testing data",
        "have people hurt themselves",
        "fda",
        "reconstitution for this exact vial",
        "0.5 ml",
        "educational information only",
        "las vegas",
        "dry (lyophilized)",
        "required documentation",
    )
    for item in products():
        reader = PdfReader(str(pdf_path(item)))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        lower = text.lower()
        for phrase in required:
            assert phrase in lower, (item["id"], phrase)
        proto = PROTOCOLS[item["id"]]
        assert f"{proto['bac_ml']:g} ml" in lower or f"{proto['bac_ml']:.1f} ml" in lower
        assert f"{proto['start_units']:g} units" in lower
        assert "not medical advice" in lower
        if item["id"] == "tirzepatide":
            assert "25" in text
            assert "2.5" in text
            assert "mounjaro" in lower
        if item["id"] == "nad":
            assert "1000" in text
            assert "50" in text
        if item["id"] == "retatrutide":
            assert "not fda-approved" in lower or "not fda approved" in lower
        if item["id"] == "ss-31":
            assert "forzinity" in lower
            assert "barth" in lower


def test_orders_workbook_covers_all_skus_and_inbox_columns():
    from openpyxl import load_workbook

    from wellness_agent.inventory.build_workbook import WORKBOOK_PATH, build_workbook

    path = build_workbook()
    assert path == WORKBOOK_PATH
    assert path.is_file()
    wb = load_workbook(path)
    assert "Inventory" in wb.sheetnames
    assert "Orders" in wb.sheetnames
    assert "Payments" in wb.sheetnames
    assert "Fulfillment" in wb.sheetnames
    assert "Shipping" in wb.sheetnames
    assert "Cancellations" in wb.sheetnames
    assert "Peptides" in wb.sheetnames
    ids = [row[0] for row in wb["Inventory"].iter_rows(min_row=2, values_only=True) if row[0]]
    assert ids == [item["id"] for item in products()]
    headers = [cell.value for cell in wb["Inventory"][1]]
    assert "on_hand" in headers


def test_ingest_files_copies_legacy_pdfs_and_workbook(tmp_path):
    from openpyxl import Workbook

    from wellness_agent.cli import main
    from wellness_agent.ingest_files import ingest_legacy
    from wellness_agent.stock import load_stock

    pdf = tmp_path / "klow.pdf"
    pdf.write_bytes((pdf_path(find_product("klow"))).read_bytes())
    extra = tmp_path / "notes.txt"
    extra.write_text("ignore", encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(["sku_id", "name", "on_hand"])
    ws.append(["klow", "KLOW", 12])
    ws.append(["semax", "Semax", None])
    orders = wb.create_sheet("Orders")
    orders.append(["order_id", "sku_id", "qty"])
    orders.append(["THW-TEST-1", "semax", 2])
    xlsx = tmp_path / "trueholdwellness-orders.xlsx"
    wb.save(xlsx)
    result = ingest_legacy(tmp_path)
    assert "klow.pdf" in result["copied_pdfs"]
    assert "notes.txt" in result["skipped"]
    assert result["workbook"] == "trueholdwellness-orders.xlsx"
    assert load_stock()["products"]["klow"]["on_hand"] == 12
    assert load_stock()["products"]["semax"]["on_hand"] is None
    assert main(["ingest-files", "--path", str(tmp_path)]) == 0


def test_picture_menu_cards_exist_for_every_sku():
    from wellness_agent.inventory.build_cards import card_path, ensure_cards

    ensure_cards()
    for item in products():
        path = card_path(item)
        assert path.is_file()
        assert path.suffix == ".jpg"
        assert path.stat().st_size > 1000


def test_brand_graphics_exist():
    from wellness_agent.inventory.build_brand import ensure_brand, hero_path, logo_path, service_path

    ensure_brand()
    for path in (logo_path(), hero_path(), service_path()):
        assert path.is_file()
        assert path.suffix.lower() in {".jpg", ".jpeg"}
        assert path.stat().st_size > 1000


def test_catalog_states_las_vegas_and_dry_vials():
    text = format_catalog().lower()
    assert "las vegas" in text
    assert "dry vials" in text
    assert "waiver" not in text
