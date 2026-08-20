from wellness_agent.cli import main
from wellness_agent.menu import parse_interest_qty
from wellness_agent.stock import (
    adjust_stock,
    format_stock,
    load_stock,
    record_order_row,
    set_on_hand,
)


def test_parse_interest_qty():
    assert parse_interest_qty("2x klow") == "2"
    assert parse_interest_qty("3 x Semax") == "3"
    assert parse_interest_qty("1 klow") == "1"
    assert parse_interest_qty("klow") == "1"
    assert parse_interest_qty("") == "1"


def test_adjust_stock_decrements_when_set():
    set_on_hand("klow", 12)
    entry = adjust_stock("klow", -3, reason="telegram-interest-order", detail="test")
    assert entry["previous"] == 12
    assert entry["on_hand"] == 9
    assert load_stock()["products"]["klow"]["on_hand"] == 9


def test_adjust_stock_keeps_unset_and_logs():
    entry = adjust_stock("semax", -2, reason="telegram-interest-order", detail="test")
    assert entry["previous"] is None
    assert entry["on_hand"] is None
    assert load_stock()["products"]["semax"]["on_hand"] is None
    assert load_stock()["adjustments"][-1]["delta"] == -2


def test_format_stock_and_cli(capsys):
    set_on_hand("ghk-cu", 7)
    text = format_stock()
    assert "TrueHold Wellness on-hand inventory" in text
    assert "GHK-Cu" in text
    assert "7 on hand" in text
    assert "not set" in text
    assert main(["stock"]) == 0
    assert "GHK-Cu" in capsys.readouterr().out


def test_record_order_row_writes_xlsx_and_on_hand(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    dest = tmp_path / "orders.xlsx"
    monkeypatch.setenv("WELLNESS_ORDERS_XLSX_PATH", str(dest))
    set_on_hand("klow", 6)
    adjust_stock("klow", -1, reason="telegram-interest-order", detail="row")
    record_order_row(
        {"id": "klow", "name": "KLOW"},
        "1",
        chat_id="88",
        phone="+17025550100",
    )
    wb = load_workbook(dest)
    assert "Orders" in wb.sheetnames
    last = list(wb["Orders"].iter_rows(min_row=2, values_only=True))[-1]
    assert "klow" in last
    assert 1 in last
    if "Inventory" in wb.sheetnames:
        header = [str(cell.value or "").strip().lower() for cell in wb["Inventory"][1]]
        if "on_hand" in header:
            sku_idx = header.index("sku_id")
            on_idx = header.index("on_hand")
            rows = {row[sku_idx]: row[on_idx] for row in wb["Inventory"].iter_rows(min_row=2, values_only=True)}
            assert rows.get("klow") == 5
