"""Ingest old-agent Finder files into the new TrueHold Wellness agent.

Drop PDFs and trueholdwellness-orders.xlsx into data/imports/legacy/, then:

    python -m wellness_agent ingest-files

https://openpyxl.readthedocs.io/
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Any

from wellness_agent.catalog import products
from wellness_agent.envfile import PACKAGE_ROOT
from wellness_agent.inventory.build_pdfs import PDF_DIR

LEGACY_DIR = PACKAGE_ROOT / "data" / "imports" / "legacy"
LEDGER_JSON = PACKAGE_ROOT / "data" / "order_ledger.json"
IMPORTED_XLSX = PACKAGE_ROOT / "data" / "imports" / "trueholdwellness-orders.xlsx"

PDF_ALIASES = {
    "tirzepatide.pdf": "tirzepatide.pdf",
    "nad-plus.pdf": "nad-plus.pdf",
    "nad.pdf": "nad-plus.pdf",
    "nad+.pdf": "nad-plus.pdf",
    "semax.pdf": "semax.pdf",
    "retatrutide.pdf": "retatrutide.pdf",
    "klow.pdf": "klow.pdf",
    "mots-c.pdf": "mots-c.pdf",
    "motsc.pdf": "mots-c.pdf",
    "ss-31.pdf": "ss-31.pdf",
    "ss31.pdf": "ss-31.pdf",
    "ghk-cu.pdf": "ghk-cu.pdf",
    "ghkcu.pdf": "ghk-cu.pdf",
}


def _normalize_name(name: str) -> str:
    return name.strip().lower().replace(" ", "-").replace("_", "-")


def map_pdf_name(filename: str) -> str | None:
    key = _normalize_name(filename)
    if key in PDF_ALIASES:
        return PDF_ALIASES[key]
    catalog = {item["pdf"].lower(): item["pdf"] for item in products()}
    return catalog.get(key)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_workbook(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    sheets: dict[str, list[dict[str, str]]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            sheets[name] = []
            continue
        headers = [_cell(item).lower().replace(" ", "_") or f"col_{idx}" for idx, item in enumerate(rows[0], start=1)]
        records = []
        for raw in rows[1:]:
            if raw is None or all(item is None or str(item).strip() == "" for item in raw):
                continue
            record = {headers[idx]: _cell(raw[idx] if idx < len(raw) else "") for idx in range(len(headers))}
            records.append(record)
        sheets[name] = records
    return {"source": str(path), "sheets": sheets}


def ingest_legacy(source_dir: Path | None = None) -> dict[str, Any]:
    folder = source_dir or LEGACY_DIR
    folder.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    copied_pdfs: list[str] = []
    skipped: list[str] = []
    workbook: dict[str, Any] | None = None
    for path in sorted(folder.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            dest_name = map_pdf_name(path.name)
            if not dest_name:
                skipped.append(path.name)
                continue
            shutil.copy2(path, PDF_DIR / dest_name)
            copied_pdfs.append(dest_name)
            continue
        if suffix in {".xlsx", ".xlsm"}:
            IMPORTED_XLSX.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, IMPORTED_XLSX)
            workbook = parse_workbook(path)
            LEDGER_JSON.write_text(json.dumps(workbook, indent=2), encoding="utf-8")
            continue
        skipped.append(path.name)
    known = {item["pdf"] for item in products()}
    missing_pdfs = sorted(name for name in known if not (PDF_DIR / name).is_file())
    return {
        "ok": True,
        "source": str(folder),
        "copied_pdfs": copied_pdfs,
        "skipped": skipped,
        "workbook": None if workbook is None else IMPORTED_XLSX.name,
        "ledger": None if workbook is None else str(LEDGER_JSON),
        "missing_pdfs": missing_pdfs,
        "catalog_pdfs": sorted(known),
    }
