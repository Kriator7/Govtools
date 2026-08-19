"""On-hand inventory for TrueHold Wellness Telegram orders.

Interest orders decrement on-hand the same way the old orders workbook did.
Starting counts come from Inventory.on_hand in the workbook (ingest-files)
or from data/stock.json. Unset counts stay unset; the adjustment is still logged.
"""

from __future__ import annotations

import json
import os
import time
from pathlib import Path
from typing import Any

from wellness_agent.catalog import products
from wellness_agent.envfile import PACKAGE_ROOT

TEMPLATE_XLSX = PACKAGE_ROOT / "wellness_agent" / "inventory" / "workbooks" / "trueholdwellness-orders.xlsx"


def stock_path() -> Path:
    override = os.environ.get("WELLNESS_STOCK_PATH")
    if override:
        return Path(override)
    return PACKAGE_ROOT / "data" / "stock.json"


def orders_xlsx_path() -> Path:
    override = os.environ.get("WELLNESS_ORDERS_XLSX_PATH")
    if override:
        return Path(override)
    return PACKAGE_ROOT / "data" / "imports" / "trueholdwellness-orders.xlsx"


def _blank() -> dict[str, Any]:
    return {
        "products": {
            item["id"]: {"name": item["name"], "on_hand": None}
            for item in products()
        },
        "adjustments": [],
    }


def load_stock() -> dict[str, Any]:
    path = stock_path()
    if not path.is_file():
        return _blank()
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return _blank()
    state = _blank()
    incoming = payload.get("products") if isinstance(payload, dict) else None
    if isinstance(incoming, dict):
        for sku, row in incoming.items():
            if sku not in state["products"] or not isinstance(row, dict):
                continue
            on_hand = row.get("on_hand")
            state["products"][sku]["on_hand"] = int(on_hand) if isinstance(on_hand, int) else None
            if isinstance(row.get("name"), str) and row["name"].strip():
                state["products"][sku]["name"] = row["name"]
    adjustments = payload.get("adjustments") if isinstance(payload, dict) else None
    if isinstance(adjustments, list):
        state["adjustments"] = adjustments[-500:]
    return state


def save_stock(state: dict[str, Any]) -> None:
    path = stock_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state, indent=2), encoding="utf-8")


def set_on_hand(sku: str, count: int) -> dict[str, Any]:
    state = load_stock()
    if sku not in state["products"]:
        raise KeyError(sku)
    state["products"][sku]["on_hand"] = int(count)
    save_stock(state)
    return dict(state["products"][sku])


def apply_on_hand_map(counts: dict[str, int]) -> dict[str, Any]:
    state = load_stock()
    for sku, count in counts.items():
        if sku in state["products"]:
            state["products"][sku]["on_hand"] = int(count)
    save_stock(state)
    return state


def adjust_stock(sku: str, delta: int, *, reason: str, detail: str = "") -> dict[str, Any]:
    state = load_stock()
    if sku not in state["products"]:
        raise KeyError(sku)
    row = state["products"][sku]
    previous = row["on_hand"]
    if isinstance(previous, int):
        row["on_hand"] = previous + int(delta)
    entry = {
        "at": int(time.time()),
        "sku": sku,
        "name": row["name"],
        "delta": int(delta),
        "previous": previous,
        "on_hand": row["on_hand"],
        "reason": reason,
        "detail": detail,
    }
    state["adjustments"].append(entry)
    state["adjustments"] = state["adjustments"][-500:]
    save_stock(state)
    return entry


def staff_inventory_line(sku: str, qty: int, *, reason: str, detail: str) -> str:
    entry = adjust_stock(sku, -abs(int(qty)), reason=reason, detail=detail)
    name = entry["name"]
    if entry["previous"] is None:
        return (
            f"Inventory: {name} logged {entry['delta']} "
            "(on-hand not set yet — put counts on the Inventory tab and ingest-files)."
        )
    return f"Inventory: {name} {entry['previous']} → {entry['on_hand']} on hand."


def format_stock() -> str:
    state = load_stock()
    lines = ["TrueHold Wellness on-hand inventory", ""]
    for item in products():
        row = state["products"].get(item["id"]) or {}
        on_hand = row.get("on_hand")
        shown = "not set" if on_hand is None else str(on_hand)
        lines.append(f"• {item['name']} — {shown} on hand ({item['vial']})")
    return "\n".join(lines) + "\n"


def record_order_row(product: dict[str, Any], qty: str, *, chat_id: str, phone: str | None) -> None:
    """Append the interest order to the working workbook, matching the old xlsx ledger."""
    from openpyxl import load_workbook

    dest = orders_xlsx_path()
    dest.parent.mkdir(parents=True, exist_ok=True)
    source = dest if dest.is_file() else TEMPLATE_XLSX
    if not source.is_file():
        return
    wb = load_workbook(source)
    if "Orders" not in wb.sheetnames:
        return
    ws = wb["Orders"]
    ws.append(
        [
            f"THW-TG-{int(time.time())}",
            time.strftime("%Y-%m-%d %H:%M"),
            "",
            phone or "",
            product["id"],
            int(qty),
            "interest",
            "Las Vegas local prep/delivery or dry-vial ship",
            "pending",
            f"Telegram chat {chat_id}",
        ]
    )
    if "Inventory" in wb.sheetnames:
        state = load_stock()
        header = [str(cell.value or "").strip().lower() for cell in wb["Inventory"][1]]
        sku_idx = header.index("sku_id") + 1 if "sku_id" in header else 1
        on_idx = header.index("on_hand") + 1 if "on_hand" in header else None
        if on_idx:
            for row in wb["Inventory"].iter_rows(min_row=2):
                sku = str(row[sku_idx - 1].value or "")
                if sku in state["products"]:
                    row[on_idx - 1].value = state["products"][sku]["on_hand"]
    wb.save(dest)
