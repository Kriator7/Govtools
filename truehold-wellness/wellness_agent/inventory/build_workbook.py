"""Build the TrueHold Wellness orders workbook (all live SKUs + inbox columns).

This is the staff ledger the old Telegram agent sent as trueholdwellness-orders.xlsx.
Drop a Finder copy into data/imports/legacy/ and run `python -m wellness_agent ingest-files`
to replace this template with the live sheet.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from wellness_agent.catalog import products

ROOT = Path(__file__).resolve().parent
WORKBOOK_DIR = ROOT / "workbooks"
WORKBOOK_PATH = WORKBOOK_DIR / "trueholdwellness-orders.xlsx"

NAVY = "092B57"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
TITLE_FONT = Font(color=NAVY, bold=True, name="Calibri", size=16)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes = "A2"


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> Worksheet:
    ws = wb.create_sheet(title)
    _header(ws, headers)
    for row in rows:
        ws.append(row)
        for col in range(1, len(headers) + 1):
            ws.cell(ws.max_row, col).alignment = WRAP
    return ws


def build_workbook(dest: Path | None = None) -> Path:
    dest = dest or WORKBOOK_PATH
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    cover = wb.active
    cover.title = "How to use"
    cover["A1"] = "TrueHold Wellness — orders & inventory ledger"
    cover["A1"].font = TITLE_FONT
    lines = [
        "",
        "This workbook is the staff ledger for @THWellness_bot.",
        "It covers every live shop SKU plus the original inbox columns:",
        "orders, payments, fulfillment, shipping, cancellations, peptides.",
        "",
        "Replace with the old-agent file from Finder:",
        "1. Telegram → Deleted Account chat → Files → Show in Finder",
        "2. Copy trueholdwellness-orders.xlsx and the PDFs",
        "3. Paste into truehold-wellness/data/imports/legacy/",
        "4. Run: python -m wellness_agent ingest-files",
        "",
        "Policy: prep and local delivery for Las Vegas residents only.",
        "Shipping is dry (lyophilized) vials only.",
        "The team calls to confirm, consult, and complete required documentation.",
    ]
    for idx, line in enumerate(lines, start=2):
        cover[f"A{idx}"] = line
        cover[f"A{idx}"].alignment = WRAP
    cover.column_dimensions["A"].width = 88

    inventory_rows = [
        [
            item["id"],
            item["name"],
            item["vial"],
            "Dry (lyophilized) vial",
            item["category"],
            item["pdf"],
            item["shop_url"],
            "Las Vegas prep/delivery; dry-vial ship only",
        ]
        for item in products()
    ]
    _sheet(
        wb,
        "Inventory",
        ["sku_id", "name", "vial", "form", "category", "pdf", "shop_url", "fulfillment"],
        inventory_rows,
    )
    _sheet(
        wb,
        "Orders",
        [
            "order_id",
            "received_at",
            "customer_name",
            "phone",
            "sku_id",
            "qty",
            "status",
            "fulfillment",
            "documentation",
            "notes",
        ],
        [
            [
                "THW-TEMPLATE-000001",
                "",
                "",
                "",
                "semax",
                1,
                "interest",
                "Las Vegas local prep/delivery or dry-vial ship",
                "pending",
                "Replace this sample row. Do not invent customer data.",
            ]
        ],
    )
    _sheet(
        wb,
        "Payments",
        ["order_id", "status", "amount", "method", "notes"],
        [],
    )
    _sheet(
        wb,
        "Fulfillment",
        ["order_id", "type", "status", "las_vegas_resident", "dry_vial_only", "notes"],
        [],
    )
    _sheet(
        wb,
        "Shipping",
        ["order_id", "carrier", "tracking", "status", "dry_vial_only", "notes"],
        [],
    )
    _sheet(
        wb,
        "Cancellations",
        ["order_id", "reason", "refund_status", "notes"],
        [],
    )
    _sheet(
        wb,
        "Peptides",
        ["received_at", "source", "sku_id", "summary", "action"],
        [],
    )
    _sheet(
        wb,
        "Clients",
        ["name", "phone", "telegram_chat_id", "las_vegas_resident", "notes"],
        [],
    )
    wb.save(dest)
    return dest


def main() -> int:
    path = build_workbook()
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
